import os
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

NIT_EMPRESA = os.getenv("EMPRESA_NIT", "No registrado")
FORMATO_MONEDA = '"$"#,##0'

def generar_excel_reporte_diario(fecha: str, ventas: list, total_general: float, total_iva: float) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de ventas"

    ws.merge_cells("A1:I1")
    ws["A1"] = "TRAZO OSCURO | REPORTE DIARIO DE VENTAS"
    ws["A1"].font = Font(size=16, bold=True, color="B91C1C")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = f"NIT: {NIT_EMPRESA}  |  Fecha consultada: {fecha}"
    ws["A2"].font = Font(size=10, italic=True, color="78716C")
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.append([])
    encabezados = ["Fecha", "N.º Venta", "Cliente", "Producto/Servicio", "Cantidad", "Precio unitario", "Subtotal ítem", "Total venta", "Estado"]
    ws.append(encabezados)
    borde = Border(bottom=Side(style="thin", color="D6D3D1"))
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = borde

    for v in ventas:
        for d in v["detalles"]:
            ws.append([
                fecha, f"#{v['id_venta']}", v["cliente_nombre"], d["nombre_item"],
                d["cantidad"], d["precio_unitario"], d["subtotal_item"], v["total"], v["estado"],
            ])

    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=9):
        for cell in row:
            cell.border = borde
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index in (6, 7, 8):
            row[index - 1].number_format = FORMATO_MONEDA

    ws.append([])
    fila_total = ws.max_row + 1
    ws.merge_cells(start_row=fila_total, start_column=6, end_row=fila_total, end_column=7)
    ws.merge_cells(start_row=fila_total, start_column=8, end_row=fila_total, end_column=9)
    ws[f"F{fila_total}"] = "TOTAL DEL DÍA"
    ws[f"H{fila_total}"] = total_general
    ws.merge_cells(start_row=fila_total + 1, start_column=6, end_row=fila_total + 1, end_column=7)
    ws.merge_cells(start_row=fila_total + 1, start_column=8, end_row=fila_total + 1, end_column=9)
    ws[f"F{fila_total + 1}"] = "IVA RECAUDADO"
    ws[f"H{fila_total + 1}"] = total_iva
    for fila in (fila_total, fila_total + 1):
        ws[f"F{fila}"].font = Font(bold=True, color="FFFFFF")
        ws[f"F{fila}"].fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
        ws[f"F{fila}"].alignment = Alignment(horizontal="right")
        ws[f"H{fila}"].font = Font(bold=True, color="FFFFFF")
        ws[f"H{fila}"].fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
        ws[f"H{fila}"].alignment = Alignment(horizontal="right")
        ws[f"H{fila}"].number_format = FORMATO_MONEDA

    for col, ancho in zip("ABCDEFGHI", [14, 12, 24, 30, 12, 16, 16, 16, 14]):
        ws.column_dimensions[col].width = ancho

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{4 + sum(len(v['detalles']) for v in ventas)}"
    ws.sheet_view.showGridLines = False

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
